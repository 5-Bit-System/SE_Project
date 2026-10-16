import io
from pathlib import Path
import docx
import openpyxl
import pytest

from app.hus.extractors import extract_attachment, extract_content
from app.hus.extractors.docx import extract_docx
from app.hus.extractors.excel import extract_excel
from app.hus.urls import CrawlError


def _create_mock_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Quy chế đào tạo", level=1)
    doc.add_paragraph("Điều 1: Quy định chung về tín chỉ.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "STT"
    table.rows[0].cells[1].text = "Môn học"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "Đại số"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _create_mock_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Học phần"
    ws.append(["Mã", "Tên", "Số tín"])
    ws.append(["MAT101", "Giải tích", 3])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_docx_extractor():
    raw_docx = _create_mock_docx()
    md = extract_docx(raw_docx)
    assert "# Quy chế đào tạo" in md
    assert "Điều 1: Quy định chung về tín chỉ." in md
    assert "| STT | Môn học |" in md
    assert "| 1 | Đại số |" in md


def test_excel_extractor():
    raw_xlsx = _create_mock_xlsx()
    md = extract_excel(raw_xlsx)
    assert "## Sheet: Học phần" in md
    assert "| Mã | Tên | Số tín |" in md
    assert "| MAT101 | Giải tích | 3 |" in md


def test_real_office_fixtures():
    root = Path(__file__).parent / "fixtures/hus"
    assert "| MAT101 | Calculus |" in extract_docx((root / "sample_table.docx").read_bytes())
    assert "| MAT101 | Calculus | 3 |" in extract_excel((root / "sample_courses.xlsx").read_bytes())


def test_extractor_factory():
    raw_docx = _create_mock_docx()
    content, needs_ocr = extract_attachment(raw_docx, "docx")
    assert not needs_ocr
    assert "Quy chế" in content

    with pytest.raises(CrawlError, match="unsupported file format"):
        extract_attachment(b"123", "exe")

    for unsupported in ("doc", "xls"):
        with pytest.raises(CrawlError, match="unsupported file format"):
            extract_attachment(b"123", unsupported)

    content, needs_ocr = extract_content(raw_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Quy chế" in content and not needs_ocr


def test_docx_keeps_paragraph_table_order():
    document = docx.Document()
    document.add_paragraph("Trước bảng")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Mã"
    table.cell(0, 1).text = "Tên"
    document.add_paragraph("Sau bảng")
    data = io.BytesIO()
    document.save(data)
    markdown = extract_docx(data.getvalue())
    assert markdown.index("Trước bảng") < markdown.index("| Mã | Tên |") < markdown.index("Sau bảng")


def test_excel_escapes_markdown_pipe():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Mã", "Tên"])
    workbook.active.append(["MAT101", "Toán | Tin"])
    data = io.BytesIO()
    workbook.save(data)
    assert "| MAT101 | Toán / Tin |" in extract_excel(data.getvalue())


def test_office_extractors_remove_zero_width_characters():
    document = docx.Document()
    document.add_paragraph("A\u200cB")
    docx_data = io.BytesIO()
    document.save(docx_data)
    assert "AB" in extract_docx(docx_data.getvalue())

    workbook = openpyxl.Workbook()
    workbook.active.append(["A\u2060B"])
    xlsx_data = io.BytesIO()
    workbook.save(xlsx_data)
    assert "AB" in extract_excel(xlsx_data.getvalue())
