"""Tests for table formatting, cell cleaning, extractor table generation, and table repair."""

from __future__ import annotations

import io
import json
from pathlib import Path
from unittest.mock import MagicMock

import docx
import openpyxl
import pytest

from app.hus.cleaning import clean_table_cell, format_markdown_table, normalize_vertical_text
from app.hus.extractors.docx import extract_docx
from app.hus.extractors.excel import extract_excel
from app.hus.extractors.pdf import _format_table as format_pdf_table
from tools.normalize_vertical_tables import normalize_all as normalize_vertical_all
from tools.repair_markdown_tables import (
    repair_all,
    repair_markdown_text,
    verify_no_broken_tables,
)


def test_clean_table_cell_variations():
    assert clean_table_cell(None) == ""
    assert clean_table_cell("") == ""
    assert clean_table_cell("hello | world") == "hello / world"
    assert clean_table_cell("line 1\r\nline 2") == "line 1<br>line 2"
    assert clean_table_cell("line 1\nline 2\nline 3") == "line 1<br>line 2<br>line 3"
    assert clean_table_cell("  <br>  line 1  <br>  ") == "line 1"
    assert clean_table_cell("line 1 <br/> line 2") == "line 1<br>line 2"
    assert clean_table_cell("<br><br>leading and trailing<br><br>") == "leading and trailing"
    assert clean_table_cell(42) == "42"


def test_format_markdown_table_single_line_guarantee():
    matrix = [
        ["Header 1\nLine 2", "Header 2\r\nLine B"],
        ["Cell 1\nwith\nmultiple\nlines", "Cell 2 | with pipe"],
        ["Normal 1", "Normal 2"],
    ]
    lines = format_markdown_table(matrix)
    assert len(lines) == 4
    for line in lines:
        assert "\n" not in line
        assert "\r" not in line
        assert line.startswith("|")
        assert line.endswith("|")

    assert lines[0] == "| Header 1<br>Line 2 | Header 2<br>Line B |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| Cell 1<br>with<br>multiple<br>lines | Cell 2 / with pipe |"
    assert lines[3] == "| Normal 1 | Normal 2 |"


def test_repair_markdown_text():
    broken = (
        "# Document Title\n\n"
        "Some introductory text.\n\n"
        "| STT | Mã\n"
        "học phần | Học phần | Số\n"
        "tín chỉ |\n"
        "| --- | --- | --- | ---\n"
        "|\n"
        "| 1 | MAT1001\n"
        "Toán 1 | Giải tích 1 | 3 |\n"
        "| 2 | MAT1002 | Đại số tuyến tính | 3 |\n\n"
        "Trailing notes.\n"
    )

    repaired, count = repair_markdown_text(broken)
    assert count == 3  # The header row, separator row, and first data row were broken

    for line in repaired.splitlines():
        s = line.strip()
        if s.startswith("|"):
            assert s.endswith("|")
            assert not (s.startswith("|") and not s.endswith("|"))

    assert "| STT | Mã<br>học phần | Học phần | Số<br>tín chỉ |" in repaired
    assert "| 1 | MAT1001<br>Toán 1 | Giải tích 1 | 3 |" in repaired
    assert "| 2 | MAT1002 | Đại số tuyến tính | 3 |" in repaired


def test_repair_markdown_text_cleans_dividers():
    broken = (
        "| Col 1 |\n"
        "Col 2 |\n"
        "| Col A <br>|\n"
        "|<br> Col B |\n"
    )
    repaired, count = repair_markdown_text(broken)
    for line in repaired.splitlines():
        assert "<br>|" not in line
        assert "|<br>" not in line


def test_docx_extractor_table_with_multiline_cells():
    doc = docx.Document()
    doc.add_heading("Khung chương trình", level=1)
    table = doc.add_table(rows=2, cols=2)

    # Cell with multiple paragraphs (which docx renders as \n)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"

    p0 = table.rows[1].cells[0].paragraphs[0]
    p0.text = "Para 1"
    p1 = table.rows[1].cells[0].add_paragraph("Para 2")
    p1.text = "Para 2"

    table.rows[1].cells[1].text = "Single para"

    bio = io.BytesIO()
    doc.save(bio)
    raw = bio.getvalue()

    md = extract_docx(raw)
    assert md
    for line in md.splitlines():
        s = line.strip()
        if s.startswith("|"):
            assert s.endswith("|")
            assert "\n" not in s

    assert "| Para 1<br>Para 2 | Single para |" in md


def test_excel_extractor_table_with_multiline_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ChuongTrinh"
    ws.append(["Mã HP", "Tên HP\n(Ghi chú)", "Số TC"])
    ws.append(["MAT1001", "Toán rời rạc\n(Toán - Tin)", 3])

    bio = io.BytesIO()
    wb.save(bio)
    raw = bio.getvalue()

    md = extract_excel(raw)
    assert md
    for line in md.splitlines():
        s = line.strip()
        if s.startswith("|"):
            assert s.endswith("|")
            assert "\n" not in s

    assert "| Mã HP | Tên HP<br>(Ghi chú) | Số TC |" in md
    assert "| MAT1001 | Toán rời rạc<br>(Toán - Tin) | 3 |" in md


def test_pdf_format_table_with_multiline_cells():
    table_data = [
        ["STT", "Học phần\n(Ngành KHDL)", "Số TC"],
        ["1", "Học máy\nvà ứng dụng", "3"],
        [None, None, None],
    ]
    lines = format_pdf_table(table_data)
    assert len(lines) == 3
    for line in lines:
        assert "\n" not in line
        assert line.startswith("|")
        assert line.endswith("|")

    assert lines[0] == "| STT | Học phần<br>(Ngành KHDL) | Số TC |"
    assert lines[2] == "| 1 | Học máy<br>và ứng dụng | 3 |"


def test_repair_all_and_verify(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "cs"
    doc_dir.mkdir(parents=True, exist_ok=True)
    file_path = doc_dir / "broken_doc--12345678.md"

    file_path.write_text(
        "| Header 1 | Header 2\n"
        "Subtitle |\n"
        "| --- | --- |\n"
        "| Value 1\n"
        "Line 2 | Value 2 |\n",
        encoding="utf-8",
    )

    # Initial check fails
    assert verify_no_broken_tables(tmp_path) == 1

    # Repair
    count, rows, paths = repair_all(tmp_path)
    assert count == 1
    assert rows == 2
    assert len(paths) == 1

    # Post-repair verify passes
    assert verify_no_broken_tables(tmp_path) == 0


def test_verify_detects_pdf_fragments_and_merged_table_widths(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "general"
    doc_dir.mkdir(parents=True)
    (doc_dir / "fragment.md").write_text(
        "| Code | Title |\n| --- | --- |\n| A | Course A |\n"
        "| Course A |\n| --- |\n| Notes |\n",
        encoding="utf-8",
    )
    (doc_dir / "merged.md").write_text(
        "| Code | Title |\n| --- | --- |\n| A | Course A |\n"
        "| Code | Title | Credits |\n| --- | --- | --- |\n| B | Course B | 3 |\n",
        encoding="utf-8",
    )

    assert verify_no_broken_tables(tmp_path) == 2


def test_verify_detects_blank_lines_between_docx_table_rows(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "computer_science"
    doc_dir.mkdir(parents=True)
    (doc_dir / "split.md").write_text(
        "| Code | Title |\n\n| --- | --- |\n\n| A | Course A |",
        encoding="utf-8",
    )

    assert verify_no_broken_tables(tmp_path) == 1


def test_verify_detects_catalog_only_broken_table(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "general"
    doc_dir.mkdir(parents=True)
    (doc_dir / "valid.md").write_text(
        "| Code | Title |\n| --- | --- |\n| A | Course A |",
        encoding="utf-8",
    )
    (tmp_path / "processed" / "catalog.json").write_text(
        json.dumps({"documents": {"abc": {"content_markdown": "| Broken |\n| --- |"}}}),
        encoding="utf-8",
    )

    assert verify_no_broken_tables(tmp_path) == 1


def test_pdf_format_table_with_vertical_cells():
    table_data = [
        ["STT", "Mã\nhọc phần", "Học phần", "ỉhc\nnít\nốS", "Số giờ tín chỉ", None, None],
        [None, None, None, None, "tếyuht\níL", "hnàh\ncựhT", "cọh\nựT"],
        ["1", "MAT1001", "Giải tích 1", "3", "30", "15", "0"],
    ]
    lines = format_pdf_table(table_data)
    assert len(lines) == 4
    assert lines[0] == "| STT | Mã<br>học phần | Học phần | Số tín chỉ | Số giờ tín chỉ |  |  |"
    assert lines[2] == "|  |  |  |  | Lý thuyết | Thực hành | Tự học |"
    for term in ("ỉhc", "tếyuht", "hnàh", "cọh"):
        for line in lines:
            assert term not in line


def test_normalize_vertical_tables_tool(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "cs"
    doc_dir.mkdir(parents=True, exist_ok=True)
    file_path = doc_dir / "sample_doc--abcdef12.md"

    sample_md = (
        "# Khung chương trình đào tạo\n\n"
        "| STT | Học phần | ỉhc<br>nít<br>ốS | tếyuht<br>íL | hnàh<br>cựhT | cọh<br>ựT |\n"
        "| --- | --- | --- | --- | --- | --- |\n"
        "| 1 | Toán 1 | 3 | 30 | 15 | 0 |\n"
    )
    file_path.write_text(sample_md, encoding="utf-8")

    catalog_path = tmp_path / "processed" / "catalog.json"
    catalog_data = {
        "version": "1.0",
        "documents": {
            "abcdef12": {
                "document_id": "abcdef12",
                "title": "Sample Doc",
                "content_markdown": sample_md,
                "content_hash": "oldhash",
            }
        },
    }
    catalog_path.write_text(json.dumps(catalog_data, ensure_ascii=False), encoding="utf-8")

    # Dry run does not modify
    file_cnt, rep_cnt, paths, cat_cnt = normalize_vertical_all(tmp_path, dry_run=True)
    assert file_cnt == 1
    assert cat_cnt == 1
    assert rep_cnt == 8  # 4 terms in md + 4 terms in catalog
    assert file_path.read_text(encoding="utf-8") == sample_md

    # Live run modifies both md and catalog.json
    file_cnt, rep_cnt, paths, cat_cnt = normalize_vertical_all(tmp_path, dry_run=False)
    assert file_cnt == 1
    assert cat_cnt == 1
    assert rep_cnt == 8

    updated_md = file_path.read_text(encoding="utf-8")
    assert "Số tín chỉ" in updated_md
    assert "Lý thuyết" in updated_md
    assert "Thực hành" in updated_md
    assert "Tự học" in updated_md
    for term in ("ỉhc", "tếyuht", "hnàh", "cọh"):
        assert term not in updated_md

    updated_catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    doc_entry = updated_catalog["documents"]["abcdef12"]
    assert "Số tín chỉ" in doc_entry["content_markdown"]
    assert doc_entry["content_hash"] != "oldhash"

    # Subsequent run finds 0 unnormalized terms
    file_cnt2, rep_cnt2, paths2, cat_cnt2 = normalize_vertical_all(tmp_path, dry_run=True)
    assert file_cnt2 == 0
    assert cat_cnt2 == 0
    assert rep_cnt2 == 0


def test_repair_markdown_tables_with_normalize_vertical(tmp_path: Path):
    doc_dir = tmp_path / "processed" / "documents" / "math"
    doc_dir.mkdir(parents=True, exist_ok=True)
    file_path = doc_dir / "broken_and_vertical--99999999.md"

    broken_vertical_content = (
        "| STT | Học phần\n"
        "chi tiết | ỉhc<br>nít<br>ốS |\n"
        "| --- | --- | --- |\n"
        "| 1 | Đại số | 3 |\n"
    )
    file_path.write_text(broken_vertical_content, encoding="utf-8")

    count, rows, paths = repair_all(tmp_path, dry_run=False, normalize_vertical=True)
    assert count == 1

    repaired_content = file_path.read_text(encoding="utf-8")
    assert "| STT | Học phần<br>chi tiết | Số tín chỉ |" in repaired_content
    assert "ỉhc" not in repaired_content
