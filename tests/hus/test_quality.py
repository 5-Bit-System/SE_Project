"""Invented source fixtures and adversarial checks for the page/cell evidence gate."""

from datetime import date
import io
from pathlib import Path

import pytest

from app.hus.extractors.html import extract_html
from app.hus.extractors.pdf import extract_pdf
from app.hus.index import HUSIndex, chunk_document
from app.hus.models import HUSDocument
from app.hus.quality import analyze_source, quality_digest, require_quality, table_from_grid
from app.hus.quality_models import QualityReview, SourceCheck, TableCheck
from app.hus.review import ReviewGateError
from tests.hus_review_fixtures import approved_review

TIME = "2026-10-26T00:00:00Z"
URL = "https://mim.hus.vnu.edu.vn/vi/synthetic-quality"
HTML = b"""<html><title>Synthetic table</title><article><h1>Synthetic table</h1>
<p>Invented test data only.</p><table><tr><th>Course</th><th>Credits</th></tr>
<tr><td>TEST001</td><td>3</td></tr></table></article></html>"""


def checked(report):
    """Explicit approvals for invented fixtures; never run on product sources."""

    def check(unit, **kwargs):
        return dict(
            outcome="verified",
            expected_sha256=quality_digest(unit),
            locator=unit.locator,
            reviewer="synthetic-reviewer",
            reviewed_at=TIME,
            **kwargs,
        )

    return QualityReview(
        report_sha256=quality_digest(report),
        page_checks={p.page_id: SourceCheck(**check(p)) for p in report.pages},
        table_checks={
            t.table_id: TableCheck(**check(t, confirmed_header_rows=t.header_rows))
            for t in report.tables
        },
        cell_checks={
            c.cell_id: SourceCheck(**check(c, observed_value=c.raw_value))
            for t in report.tables
            for c in t.cells
            if c.presence != "covered"
        },
    )


def table_fixture():
    doc, _ = extract_html(HTML, URL, TIME)
    report = analyze_source(HTML, URL, TIME)
    review = approved_review(doc, HTML)
    review.source_quality = checked(report)
    return doc, report, review


def test_ragged_grid_preserves_missing_empty_and_raw_values():
    rows = [["Name", "Value"], ["  A\nB  ", ""], ["Only one cell"]]
    table = table_from_grid("test", rows, {"page": 1})
    assert table.raw_rows == rows and len(table.raw_rows[-1]) == 1
    assert table.cells[2].raw_value == "  A\nB  "
    assert table.cells[2].display_value != table.cells[2].raw_value
    assert table.cells[3].presence == "empty"
    assert table.cells[-1].presence == "missing"
    assert len({c.cell_id for c in table.cells}) == 6
    assert table.cells[-1].locator == {"page": 1, "row": 2, "column": 1}


def test_html_multilevel_headers_and_merged_geometry():
    raw = b"""<article><table><tr><th rowspan='2'>Code</th><th colspan='2'>Credits</th></tr>
    <tr><th>Required</th><th>Optional</th></tr><tr><td>A</td><td>3</td><td></td></tr></table></article>"""
    table = analyze_source(raw, URL, TIME).tables[0]
    assert table.header_rows == [0, 1] and table.header_status == "explicit"
    assert table.cells[0].rowspan == 2 and table.cells[1].colspan == 2
    assert table.cells[2].covered_by == table.cells[1].cell_id
    assert table.cells[3].covered_by == table.cells[0].cell_id
    assert table.header_paths[table.column_ids[2]] == [
        table.cells[1].cell_id,
        table.cells[5].cell_id,
    ]
    assert table.cells[-1].presence == "empty" and table.raw_rows[0][2] is None
    assert table.locator["dom_path"] == "table[1]"


def test_html_row_headers_do_not_turn_every_data_row_into_column_header():
    raw = b"<article><table><tr><th>Code</th><th>Value</th></tr><tr><th scope='row'>A</th><td>3</td></tr></table></article>"
    table = analyze_source(raw, URL, TIME).tables[0]
    assert table.header_rows == [0]
    assert table.cells[2].role == "header" and table.cells[2].header_scope == "row"
    assert table.cells[3].role == "data"


def test_html_overlapping_spans_are_not_silently_repaired():
    raw = b"<article><table><tr><td>A</td><td rowspan='2'>B</td></tr><tr><td colspan='2'>C</td></tr></table></article>"
    with pytest.raises(ValueError, match="overlapping"):
        analyze_source(raw, URL, TIME)


def test_docx_preserves_merged_cells_blank_rows_and_block_locator():
    import docx
    from docx.oxml import OxmlElement

    document = docx.Document()
    document.add_paragraph("Invented source")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Two-column heading"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).text, table.cell(1, 1).text = "A", "3"
    buf = io.BytesIO()
    document.save(buf)
    result = analyze_source(buf.getvalue(), URL + ".docx", TIME).tables[0]
    assert result.locator["block"] == 1 and result.header_rows == [0]
    assert len(result.raw_rows) == 3 and result.raw_rows[-1] == ["", ""]
    assert result.cells[0].colspan == 2 and result.cells[1].presence == "covered"


def test_xlsx_keeps_formula_separate_from_evaluated_value_and_merge():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Credits", None])
    sheet.merge_cells("A1:B1")
    sheet.append([3, "=A2+1"])
    buf = io.BytesIO()
    workbook.save(buf)
    report = analyze_source(buf.getvalue(), URL + ".xlsx", TIME)
    table = report.tables[0]
    assert table.cells[0].colspan == 2 and table.cells[1].presence == "covered"
    assert table.cells[-1].raw_value == "=A2+1" and table.cells[-1].presence == "unresolved"
    assert table.cells[-1].locator["address"] == "B2"
    with pytest.raises(ValueError, match="unresolved"):
        require_quality(report, checked(report))


def test_mixed_pdf_is_flagged_per_page_instead_of_document_average(monkeypatch):
    class Page:
        images, lines, rects, curves = [], [], [], []

        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

        def find_tables(self):
            return []

    class PDF:
        pages = [Page("Long text " * 100), Page(""), Page("Last text " * 100)]

        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

    PDF.pages[1].images = [{"name": "scanned-page"}]
    monkeypatch.setattr("pdfplumber.open", lambda _: PDF())
    assert extract_pdf(b"%PDF-mock")[1] is True
    report = analyze_source(b"%PDF-mock", URL + ".pdf", TIME)
    assert [page.page_number for page in report.pages if page.needs_ocr] == [2]
    with pytest.raises(ValueError, match="OCR/visual"):
        require_quality(report, checked(report))


def test_blank_pdf_requires_explicit_visual_resolution():
    import pypdfium2

    pdf = pypdfium2.PdfDocument.new()
    pdf.new_page(200, 200)
    buf = io.BytesIO()
    pdf.save(buf)
    pdf.close()
    report = analyze_source(buf.getvalue(), URL + ".pdf", TIME)
    checks = checked(report)
    with pytest.raises(ValueError, match="OCR/visual"):
        require_quality(report, checks)
    checks.page_checks[report.pages[0].page_id].resolution = "blank_verified"
    require_quality(report, checks)


def test_real_pdf_structure_and_fully_checked_positive_path():
    raw = (Path(__file__).parents[1] / "fixtures/hus/sample_text.pdf").read_bytes()
    report = analyze_source(raw, URL + ".pdf", TIME)
    table = report.tables[0]
    assert table.cells[0].locator["page"] == 1 and len(table.cells[0].locator["bbox"]) == 4
    assert any(cell.raw_value == "MAT101" for cell in table.cells)
    doc = HUSDocument(
        source_url=URL + ".pdf",
        title="Invented test PDF",
        content_markdown=report.display_content,
        content_hash=report.display_content_sha256,
        fetched_at=TIME,
    )
    decision = approved_review(doc, raw)
    decision.source_quality = checked(report)
    chunks = chunk_document(doc, review=decision, raw=raw)
    assert chunks and chunks[0].source_quality_sha256 == quality_digest(report)
    assert report.pages[0].page_id in chunks[0].source_unit_ids
    decision.source_quality.page_checks.clear()
    with pytest.raises(ReviewGateError, match="exactly"):
        chunk_document(doc, review=decision, raw=raw)


@pytest.mark.parametrize(
    "change", ["none", "table", "cell", "locator", "value", "hash", "extra", "pending", "future"]
)
def test_table_review_rejects_incomplete_stale_or_false_checks(change):
    doc, report, review = table_fixture()
    checks = review.source_quality
    last = report.tables[0].cells[-1].cell_id
    if change == "none":
        review.source_quality = None
    elif change == "table":
        checks.table_checks.clear()
    elif change == "cell":
        checks.cell_checks.pop(last)
    elif change == "locator":
        checks.cell_checks[last].locator = {"row": 999}
    elif change == "value":
        checks.cell_checks[last].observed_value = "30"
    elif change == "hash":
        checks.report_sha256 = "0" * 64
    elif change == "extra":
        checks.cell_checks["foreign"] = checks.cell_checks[last]
    elif change == "pending":
        checks.cell_checks[last].outcome = "unresolved"
    elif change == "future":
        checks.cell_checks[last].reviewed_at = "2026-10-27T00:00:00Z"
    with pytest.raises(ReviewGateError):
        chunk_document(doc, review=review, raw=HTML, as_of=date(2026, 10, 26))


def test_modified_display_numbers_cannot_reuse_source_checks():
    import hashlib

    doc, _, review = table_fixture()
    doc.content_markdown = doc.content_markdown.replace("| 3 |", "| 30 |")
    doc.content_hash = hashlib.sha256(doc.content.encode()).hexdigest()
    review.content_hash = doc.content_hash
    with pytest.raises(ReviewGateError, match="text differs"):
        chunk_document(doc, review=review, raw=HTML)


def test_ragged_html_cannot_be_approved_by_padded_display():
    raw = b"<article><table><tr><th>A</th><th>B</th></tr><tr><td>3</td></tr></table></article>"
    doc, _ = extract_html(raw, URL, TIME)
    report = analyze_source(raw, URL, TIME)
    decision = approved_review(doc, raw)
    decision.source_quality = checked(report)
    assert report.tables[0].cells[-1].presence == "unresolved"
    with pytest.raises(ReviewGateError, match="unresolved"):
        chunk_document(doc, review=decision, raw=raw)


def test_index_rechecks_live_cell_decisions_and_drops_cached_approval(tmp_path):
    doc, report, decision = table_fixture()
    reviews = {doc.url: decision}
    index = HUSIndex(tmp_path / "index.json", tmp_path / "raw", reviews=reviews)
    index.add(doc, HTML)
    index.save()
    assert index.documents[doc.url]["source_quality"]["source_sha256"] == report.source_sha256
    assert HUSIndex(index.path, index.raw_dir, reviews=reviews).documents
    decision.source_quality.cell_checks.clear()
    assert HUSIndex(index.path, index.raw_dir, reviews=reviews).documents == {}
    index.save()
    assert index.documents == {}


def test_parser_failure_is_a_gate_error(monkeypatch):
    doc, _, review = table_fixture()

    def fail(*_):
        raise RuntimeError("broken source parser")

    monkeypatch.setattr("app.hus.quality.analyze_source", fail)
    with pytest.raises(ReviewGateError, match="broken source parser"):
        chunk_document(doc, review=review, raw=HTML)


def test_index_drops_record_if_report_cannot_be_recreated_after_gate(tmp_path, monkeypatch):
    doc, _, decision = table_fixture()
    reviews = {doc.url: decision}
    index = HUSIndex(tmp_path / "index.json", tmp_path / "raw", reviews=reviews)
    index.add(doc, HTML)
    index.save()
    calls = 0

    def unstable(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("source report unavailable")
        return analyze_source(*args, **kwargs)

    monkeypatch.setattr("app.hus.quality.analyze_source", unstable)
    assert HUSIndex(index.path, index.raw_dir, reviews=reviews).documents == {}


def synthetic_snapshot(tmp_path, raw, url):
    import hashlib
    from app.hus.corpus import create_source_manifest, initial_reviews, json_bytes

    root = tmp_path / "sources"
    key = "a" * 16
    report = analyze_source(raw, url, TIME)
    doc = HUSDocument(
        document_id=key,
        source_url=url,
        title="Invented source",
        content_markdown=report.display_content,
        content_hash=report.display_content_sha256,
        fetched_at=TIME,
    )
    markdown = root / f"processed/documents/general/fixture--{key[:8]}.md"
    markdown.parent.mkdir(parents=True)
    markdown.write_text(f"---\ntitle: Invented source\n---\n\n{doc.content}", encoding="utf-8")
    suffix = ".pdf" if url.endswith(".pdf") else ".html"
    raw_path = (
        root
        / f"raw/{'documents' if suffix == '.pdf' else 'html'}/fixture--{hashlib.sha1(url.encode()).hexdigest()[:12]}{suffix}"
    )
    raw_path.parent.mkdir(parents=True)
    raw_path.write_bytes(raw)
    (root / "processed/catalog.json").write_bytes(
        json_bytes({"document_count": 1, "documents": {key: doc.model_dump()}})
    )
    manifest = create_source_manifest(root, "synthetic-quality", "fixture://synthetic-quality")
    reviews = initial_reviews(root, manifest)
    return root, manifest, reviews, doc, report


def test_corpus_cannot_export_unchecked_cell_and_keeps_source_references(tmp_path):
    import json
    from app.hus.corpus import BuildConfig, CorpusError, build_artifacts
    from app.hus.review import source_digest

    root, manifest, reviews, doc, report = synthetic_snapshot(tmp_path, HTML, URL)
    review = approved_review(doc, HTML)
    review.source_checksum_encoding = "utf8-lf"
    review.reviewed_source_sha256 = source_digest(HTML, "utf8-lf")
    review.source_quality = checked(report)
    reviews.decisions[doc.document_id] = review
    config = BuildConfig(
        schema_version="1.0",
        pipeline_version="1.4",
        source_manifest="source_manifest.json",
        review_decisions="reviews.json",
        normalization="nfc-lf-trim-trailing-v1",
        chunking="utf8-lines-v1",
        max_chunk_bytes=1200,
        overlap_bytes=0,
    )
    # A checksum of binary HTML cannot stand in for the LF-normalized source report.
    with pytest.raises(CorpusError, match="report hash"):
        build_artifacts(root, manifest, reviews, config)
    report = analyze_source(HTML, URL, TIME, "utf8-lf")
    review.source_quality = checked(report)
    artifacts = build_artifacts(root, manifest, reviews, config)
    chunk = json.loads(artifacts["chunks.jsonl"].splitlines()[0])
    assert chunk["source_quality_sha256"] == quality_digest(report)
    assert report.tables[0].table_id in chunk["source_unit_ids"]
    assert (
        json.loads(artifacts["source_map.jsonl"].splitlines()[0])["quality_scope"]
        == "whole_document"
    )
    review.source_quality.cell_checks.pop(report.tables[0].cells[-1].cell_id)
    with pytest.raises(CorpusError, match="exactly"):
        build_artifacts(root, manifest, reviews, config)


def test_inspector_creates_page_ocr_queue_and_never_approves(tmp_path):
    import json
    import pypdfium2
    from app.hus.corpus import json_bytes
    from tools.inspect_hus_source_quality import inspect

    pdf = pypdfium2.PdfDocument.new()
    pdf.new_page(200, 200)
    buf = io.BytesIO()
    pdf.save(buf)
    pdf.close()
    root, manifest, _, doc, _ = synthetic_snapshot(tmp_path, buf.getvalue(), URL + ".pdf")
    contract = tmp_path / "source_manifest.json"
    contract.write_bytes(json_bytes(manifest.model_dump()))
    out = tmp_path / "inspection"
    summary = inspect(root, contract, None, out)
    assert summary["pdf_pages"] == summary["ocr_pages"] == 1
    assert summary["verified_pages"] == summary["verified_cells"] == 0
    queue = json.loads((out / "ocr_queue.jsonl").read_text())
    assert queue["document_id"] == doc.document_id and queue["page"] == 1
    assert queue["status"] == "needs_visual_or_ocr_review"
    assert not (out / "review_decisions.json").exists()
    relocated = tmp_path / "inspection2"
    inspect(root, contract, None, relocated)
    assert {p.name: p.read_bytes() for p in out.iterdir()} == {
        p.name: p.read_bytes() for p in relocated.iterdir()
    }
    with pytest.raises(ValueError, match="already exists"):
        inspect(root, contract, None, out)
