"""Re-extract byte-bound page/table evidence and fail closed on unchecked content."""

from datetime import date, datetime
import hashlib
import io
import json
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
import zipfile

from app.hus.cleaning import clean_table_cell, clean_text, format_markdown_table
from app.hus.quality_models import (
    PageQuality,
    QualityReview,
    SourceCell,
    SourceCheck,
    SourceQuality,
    SourceTable,
)


def quality_digest(value: Any) -> str:
    if hasattr(value, "model_dump"):
        value = value.model_dump()
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def kind_for(raw: bytes, url: str) -> str:
    if raw.lstrip().startswith(b"%PDF"):
        return "pdf"
    if raw.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if "word/document.xml" in archive.namelist():
                return "docx"
            if "xl/workbook.xml" in archive.namelist():
                return "xlsx"
    suffix = Path(urlsplit(url).path).suffix.lower().lstrip(".")
    if suffix in {"pdf", "docx", "xlsx"}:
        return suffix
    if b"<" in raw[:1024]:
        return "html"
    return "unknown"


def table_from_grid(
    table_id: str,
    raw_rows: list[list[str | None]],
    locator: dict,
    *,
    geometry: dict[tuple[int, int], tuple[int, int, int, int, list[float] | None]] | None = None,
    header_rows: list[int] | None = None,
    explicit_headers: bool = False,
    header_scopes: dict[tuple[int, int], str] | None = None,
) -> SourceTable:
    """Keep missing, empty, and covered slots distinct; never pad the raw rows."""
    width = max((len(row) for row in raw_rows), default=0)
    row_ids = [f"{table_id}:r{r}" for r in range(len(raw_rows))]
    columns = [f"{table_id}:c{c}" for c in range(width)]
    cells, issues = [], set()
    for r, row in enumerate(raw_rows):
        for c in range(width):
            missing = c >= len(row)
            value = None if missing else row[c]
            origin = geometry.get((r, c)) if geometry is not None else None
            covered = origin is not None and origin[:2] != (r, c)
            presence = (
                "missing"
                if missing
                else "covered"
                if covered
                else "unresolved"
                if value is None and geometry is not None and origin is None
                else "empty"
                if value in (None, "")
                else "present"
            )
            if presence in {"missing", "unresolved"}:
                issues.add("unresolved_grid_position")
            rr, cc, rs, cs, bbox = origin if origin else (r, c, 1, 1, None)
            if origin and (rr + rs > len(raw_rows) or cc + cs > width):
                issues.add("span_outside_grid")
            if covered and locator.get("kind") == "pdf" and value not in (None, ""):
                issues.add("content_in_covered_grid_position")
            cell_locator = {**locator, "row": r, "column": c}
            if bbox is not None:
                cell_locator["bbox"] = bbox
            cells.append(
                SourceCell(
                    cell_id=f"{table_id}:r{r}:c{c}",
                    row_id=row_ids[r],
                    column_id=columns[c],
                    row=r,
                    column=c,
                    raw_value=value,
                    display_value=None if covered else clean_table_cell(value),
                    presence=presence,
                    rowspan=rs if not covered else 1,
                    colspan=cs if not covered else 1,
                    covered_by=f"{table_id}:r{rr}:c{cc}" if covered else None,
                    locator=cell_locator,
                )
            )
    headers = header_rows if header_rows is not None else ([0] if raw_rows else [])
    for cell in cells:
        scope = (header_scopes or {}).get((cell.row, cell.column))
        cell.role = "header" if scope or cell.row in headers else "data"
        cell.header_scope = scope or ("column" if cell.row in headers else "none")
    paths = {}
    for c, column in enumerate(columns):
        ids = []
        for cell in cells:
            if cell.row in headers and cell.column == c:
                key = cell.covered_by or cell.cell_id
                if key not in ids:
                    ids.append(key)
        paths[column] = ids
    return SourceTable(
        table_id=table_id,
        locator=locator,
        row_ids=row_ids,
        column_ids=columns,
        raw_rows=raw_rows,
        cells=cells,
        header_rows=headers,
        header_status="explicit" if explicit_headers else "proposed" if headers else "unknown",
        header_paths=paths,
        display_markdown="\n".join(format_markdown_table(raw_rows)),
        issues=sorted(issues),
    )


def _pdf_tables(page, page_no: int, prefix: str) -> list[SourceTable]:
    tables = []
    for index, table in enumerate(sorted(page.find_tables(), key=lambda item: item.bbox[1])):
        rows = table.extract()
        if not rows or max(map(len, rows)) < 2:
            continue
        raw_rows = [[None if value is None else str(value) for value in row] for row in rows]
        pdf_rows = getattr(table, "rows", [])
        boxes = [box for row in pdf_rows for box in row.cells if box]
        xs = sorted({round(float(x), 3) for box in boxes for x in (box[0], box[2])})
        ys = sorted({round(float(y), 3) for box in boxes for y in (box[1], box[3])})
        geometry = {}
        for row in pdf_rows:
            for box in row.cells:
                if not box:
                    continue
                x0, top, x1, bottom = [round(float(v), 3) for v in box]
                c0, c1, r0, r1 = xs.index(x0), xs.index(x1), ys.index(top), ys.index(bottom)
                for r in range(r0, r1):
                    for c in range(c0, c1):
                        geometry[(r, c)] = (r0, c0, r1 - r0, c1 - c0, [x0, top, x1, bottom])
        table_id = f"{prefix}:p{page_no}:t{index}"
        tables.append(
            table_from_grid(
                table_id,
                raw_rows,
                {
                    "kind": "pdf",
                    "page": page_no,
                    "table": index,
                    "bbox": [round(float(v), 3) for v in table.bbox],
                },
                geometry=geometry,
            )
        )
    return tables


def _docx_tables(raw: bytes, prefix: str) -> list[SourceTable]:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(raw))
    result = []
    for block_no, block in enumerate(document.iter_inner_content()):
        if not isinstance(block, Table):
            continue
        rows, positions, headers = [], {}, []
        for r, row in enumerate(block.rows):
            offset = getattr(row, "grid_cols_before", 0)
            values = [None] * offset
            for c, cell in enumerate(row.cells, start=offset):
                values.append(cell.text)
                positions.setdefault(cell._tc, []).append((r, c))
            rows.append(values)
            if (
                row._tr.trPr is not None
                and row._tr.trPr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
                )
                is not None
            ):
                headers.append(r)
        geometry = {}
        for coordinates in positions.values():
            r0, c0 = min(r for r, _ in coordinates), min(c for _, c in coordinates)
            rs = max(r for r, _ in coordinates) - r0 + 1
            cs = max(c for _, c in coordinates) - c0 + 1
            for r, c in coordinates:
                geometry[(r, c)] = (r0, c0, rs, cs, None)
        table_id = f"{prefix}:b{block_no}"
        result.append(
            table_from_grid(
                table_id,
                rows,
                {
                    "kind": "docx",
                    "block": block_no,
                    "xpath": f"/w:document/w:body/*[{block_no + 1}]",
                },
                geometry=geometry,
                header_rows=headers or None,
                explicit_headers=bool(headers),
            )
        )
    return result


def _html_tables(raw: bytes, prefix: str) -> list[SourceTable]:
    from bs4 import BeautifulSoup
    from app.hus.extractors.html import _DROP, _main

    soup = BeautifulSoup(raw, "html.parser", from_encoding="utf-8")
    original_indexes = {id(table): index for index, table in enumerate(soup.find_all("table"))}
    for tag in soup.select(_DROP):
        tag.decompose()
    result = []
    for table in _main(soup).find_all("table"):
        index = original_indexes[id(table)]
        values, geometry, headers, scopes = {}, {}, [], {}
        trs = [row for row in table.find_all("tr") if row.find_parent("table") is table]
        for r, row in enumerate(trs):
            c = 0
            for cell in row.find_all(["td", "th"], recursive=False):
                while (r, c) in values:
                    c += 1
                rs, cs = max(1, int(cell.get("rowspan", 1))), max(1, int(cell.get("colspan", 1)))
                if rs > 128 or cs > 128 or r + rs > len(trs):
                    raise ValueError("Source table span exceeds supported bound")
                value = cell.get_text("\n", strip=False)
                for rr in range(r, min(r + rs, len(trs))):
                    for cc in range(c, c + cs):
                        if (rr, cc) in values:
                            raise ValueError("Source table contains overlapping merged cells")
                        values[(rr, cc)] = value if (rr, cc) == (r, c) else None
                        geometry[(rr, cc)] = (r, c, rs, cs, None)
                if cell.name == "th":
                    explicit_scope = cell.get("scope", "")
                    scope = {
                        "col": "column",
                        "row": "row",
                        "colgroup": "column_group",
                        "rowgroup": "row_group",
                    }.get(explicit_scope, "unknown")
                    if scope == "unknown" and (
                        row.find_parent("thead")
                        or all(
                            item.name == "th"
                            for item in row.find_all(["td", "th"], recursive=False)
                        )
                    ):
                        scope = "column"
                    scopes[(r, c)] = scope
                    if scope in {"column", "column_group"} and r not in headers:
                        headers.append(r)
                c += cs
        width = max((c + 1 for _, c in values), default=0)
        rows = [[values.get((r, c)) for c in range(width)] for r in range(len(trs))]
        result.append(
            table_from_grid(
                f"{prefix}:t{index}",
                rows,
                {"kind": "html", "dom_path": f"table[{index + 1}]", "table": index},
                geometry=geometry,
                header_rows=headers,
                explicit_headers=bool(headers),
                header_scopes=scopes,
            )
        )
    return result


def _xlsx_tables(raw: bytes, prefix: str) -> list[SourceTable]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    result = []
    try:
        for index, sheet in enumerate(workbook):
            rows, geometry = [], {}
            for r, row in enumerate(sheet.iter_rows()):
                rows.append([None if cell.value is None else str(cell.value) for cell in row])
                for c, cell in enumerate(row):
                    geometry[(r, c)] = (r, c, 1, 1, None)
            for span in sheet.merged_cells.ranges:
                for r in range(span.min_row - 1, span.max_row):
                    for c in range(span.min_col - 1, span.max_col):
                        geometry[(r, c)] = (
                            span.min_row - 1,
                            span.min_col - 1,
                            span.max_row - span.min_row + 1,
                            span.max_col - span.min_col + 1,
                            None,
                        )
            table = table_from_grid(
                f"{prefix}:s{index}",
                rows,
                {"kind": "xlsx", "sheet": sheet.title},
                geometry=geometry,
            )
            for cell in table.cells:
                cell.locator["address"] = sheet.cell(cell.row + 1, cell.column + 1).coordinate
                if sheet.cell(cell.row + 1, cell.column + 1).data_type == "f":
                    cell.presence = "unresolved"
                    table.issues.append("formula_value_requires_reviewed_source_export")
            result.append(table)
    finally:
        workbook.close()
    return result


def analyze_source(
    raw: bytes, source_url: str, fetched_at: str, encoding: str = "binary"
) -> SourceQuality:
    if encoding == "utf8-lf":
        raw = raw.decode("utf-8").replace("\r\n", "\n").replace("\r", "\n").encode()
    elif encoding != "binary":
        raise ValueError("Unknown quality source encoding")
    digest = hashlib.sha256(raw).hexdigest()
    prefix, kind = digest[:16], kind_for(raw, source_url)
    pages, tables, content = [], [], ""
    if kind == "pdf":
        import pdfplumber
        from app.hus.extractors.pdf import extract_pdf

        content, _ = extract_pdf(raw)
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for index, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                chars = len(clean_text(text))
                pages.append(
                    PageQuality(
                        page_id=f"{prefix}:p{index}",
                        page_number=index,
                        text_chars=chars,
                        image_count=len(getattr(page, "images", [])),
                        graphic_count=sum(
                            len(getattr(page, name, [])) for name in ("lines", "rects", "curves")
                        ),
                        needs_ocr=chars < 50,
                        reasons=["low_text_requires_visual_review"] if chars < 50 else [],
                        raw_text=text,
                        display_text=clean_text(text),
                        locator={"kind": "pdf", "page": index},
                    )
                )
                tables.extend(_pdf_tables(page, index, prefix))
    elif kind == "docx":
        from app.hus.extractors.docx import extract_docx

        content, tables = extract_docx(raw), _docx_tables(raw, prefix)
    elif kind == "xlsx":
        from app.hus.extractors.excel import extract_excel

        content, tables = extract_excel(raw), _xlsx_tables(raw, prefix)
    elif kind == "html":
        from app.hus.extractors.html import extract_html

        doc, _ = extract_html(raw, source_url, fetched_at)
        content, tables = doc.content, _html_tables(raw, prefix)
    return SourceQuality(
        source_url=source_url,
        source_sha256=digest,
        checksum_encoding=encoding,
        source_kind=kind,
        display_content=content,
        display_content_sha256=hashlib.sha256(content.strip().encode()).hexdigest(),
        pages=pages,
        tables=tables,
        issues=["unsupported_source"] if kind == "unknown" else [],
    )


def _check(target, review: SourceCheck | None, as_of: date | None) -> None:
    if review is None or review.outcome != "verified":
        raise ValueError("Source page/table/cell has not been verified")
    if review.expected_sha256 != quality_digest(target) or review.locator != target.locator:
        raise ValueError("Source check is stale or points to a different locator")
    if as_of and datetime.fromisoformat(review.reviewed_at.replace("Z", "+00:00")).date() > as_of:
        raise ValueError("Source quality review is later than as_of")


def require_quality(
    report: SourceQuality, review: QualityReview | None, as_of: date | None = None
) -> None:
    if not report.pages and not report.tables:
        return
    if review is None:
        raise ValueError("Missing page/table quality review")
    review = QualityReview.model_validate(review.model_dump())
    if review.report_sha256 != quality_digest(report):
        raise ValueError("Source quality report hash has changed")
    if (
        set(review.page_checks) != {page.page_id for page in report.pages}
        or set(review.table_checks) != {table.table_id for table in report.tables}
        or set(review.cell_checks)
        != {
            cell.cell_id
            for table in report.tables
            for cell in table.cells
            if cell.presence != "covered"
        }
    ):
        raise ValueError("Source quality checks must cover exactly the source page/table/cell IDs")
    for page in report.pages:
        check = review.page_checks.get(page.page_id)
        _check(page, check, as_of)
        if check.resolution == "ocr_pending":
            raise ValueError("Source page still awaits OCR")
        if page.needs_ocr:
            blank = not page.raw_text.strip() and page.image_count == page.graphic_count == 0
            short = bool(page.raw_text.strip()) and page.image_count == 0
            if not (
                (blank and check.resolution == "blank_verified")
                or (short and check.resolution == "short_text_verified")
            ):
                raise ValueError("OCR/visual review is unresolved for a source page")
    for table in report.tables:
        if table.issues:
            raise ValueError("Source table has unresolved structural/value issues")
        check = review.table_checks.get(table.table_id)
        _check(table, check, as_of)
        if len(set(check.confirmed_header_rows)) != len(check.confirmed_header_rows) or any(
            row < 0 or row >= len(table.row_ids) for row in check.confirmed_header_rows
        ):
            raise ValueError("Invalid reviewed header rows")
        for cell in table.cells:
            if cell.presence == "covered":
                continue
            if cell.presence in {"missing", "unresolved"}:
                raise ValueError("Missing/unresolved source cell cannot be published")
            cell_check = review.cell_checks.get(cell.cell_id)
            _check(cell, cell_check, as_of)
            if cell_check.observed_value != cell.raw_value:
                raise ValueError("Reviewed cell value differs from source extraction")
