"""XLSX spreadsheet extraction using openpyxl."""

from __future__ import annotations

import io

import openpyxl

from app.hus.cleaning import format_markdown_table
from app.hus.urls import CrawlError


def extract_excel(raw: bytes) -> str:
    """Extract sheets and table data from XLSX bytes to markdown tables."""
    if not raw:
        raise CrawlError("empty xlsx file")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sections: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
            table_lines = format_markdown_table(raw_rows)
            if not table_lines:
                continue

            sheet_lines = [f"## Sheet: {sheet_name}", *table_lines]
            sections.append("\n".join(sheet_lines))

        content = "\n\n".join(sections).strip()
        if not content:
            raise CrawlError("xlsx contains no readable tabular data")
        return content
    except Exception as exc:
        if isinstance(exc, CrawlError):
            raise
        raise CrawlError(f"excel extraction failed: {exc}") from exc
